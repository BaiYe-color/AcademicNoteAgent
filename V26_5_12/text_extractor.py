#!/usr/bin/env python3
"""PDF 和 PPTX 文本提取"""
from pathlib import Path
import pdfplumber
from pptx import Presentation
from utils import console, create_progress, print_status, print_warning

def extract_text(pdf_path: str) -> str:
    print_status("正在提取 PDF 文本...", style="cyan", icon="📄")
    path = Path(pdf_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"找不到 PDF 文件：{path}")

    all_pages = []
    with pdfplumber.open(path) as pdf:
        total = len(pdf.pages)
        with create_progress() as progress:
            task = progress.add_task("[cyan]Extracting PDF…", total=total)
            for i, page in enumerate(pdf.pages, start=1):
                text = page.extract_text()
                if text:
                    all_pages.append(text.strip())
                else:
                    all_pages.append(f"[警告] 第{i}页无可提取文本，可能为扫描图片，请自行处理。")
                    print_warning(f"第{i}页无可提取文本")
                progress.update(task, advance=1)
    print_status(f"文本提取完成，共 {total} 页。", style="green", icon="✔")
    return "\n\n".join(all_pages)

def extract_text_from_pptx(pptx_path: str) -> str:
    print_status("正在提取 PPT 文本...", style="cyan", icon="📊")
    path = Path(pptx_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"找不到 PPT 文件：{path}")

    prs = Presentation(str(path))
    results = []
    slide_count = len(prs.slides)
    with create_progress() as progress:
        task = progress.add_task("[cyan]Extracting PPTX…", total=slide_count)
        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        txt = para.text.strip()
                        if txt:
                            slide_texts.append(txt)
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        for cell in row.cells:
                            txt = cell.text.strip()
                            if txt:
                                slide_texts.append(txt)
            if slide_texts:
                slide_content = "\n".join(slide_texts)
            else:
                slide_content = f"[警告] 第{slide_num}页没有文字"
                print_warning(f"第{slide_num}页没有文字")
            results.append(f"--- Slide {slide_num} ---\n{slide_content}")
            progress.update(task, advance=1)
    print_status(f"文本提取完成，共 {slide_count} 页。", style="green", icon="✔")
    return "\n\n".join(results)